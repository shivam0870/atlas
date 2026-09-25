"""Bounded, offline document extraction with stable page/section offsets."""

import asyncio
import base64
import io
import json
import os
import sys
import tempfile
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path
from xml.etree import ElementTree

from atlas.config import settings

MAX_UPLOAD_BYTES = settings.max_file_bytes
MAX_EXPANDED_BYTES = 20 * 1024 * 1024
MAX_CHARACTERS = 1_000_000
MAX_PAGES = settings.max_document_pages
SUPPORTED = {".txt", ".md", ".pdf", ".docx"}
WORKER_SCRIPT = str(Path(__file__).resolve())


class ExtractionError(ValueError):
    pass


@dataclass(frozen=True)
class ExtractedDocument:
    content: str
    media_type: str
    segments: list[dict]


def _join(parts: list[tuple[str, dict]], media_type: str) -> ExtractedDocument:
    text = ""
    segments = []
    for value, attributes in parts:
        value = value.replace("\x00", "").strip()
        if not value:
            continue
        if text:
            text += "\n\n"
        start = len(text)
        text += value
        if len(text) > MAX_CHARACTERS:
            raise ExtractionError("Extracted document exceeds the one-million-character limit")
        segments.append({"start": start, "end": len(text), **attributes})
    if not text:
        raise ExtractionError("Document has no readable text")
    return ExtractedDocument(text, media_type, segments)


def extract_bytes(filename: str, data: bytes) -> ExtractedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise ExtractionError("Supported file types are TXT, Markdown, text PDF and DOCX")
    if not data or len(data) > MAX_UPLOAD_BYTES:
        raise ExtractionError("Upload a nonempty file no larger than 20 MB")
    if suffix in {".txt", ".md"}:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ExtractionError("Text files must use UTF-8 encoding") from exc
        if "\x00" in text:
            raise ExtractionError("This appears to be a binary file, not UTF-8 text")
        return _join(
            [(text, {"section": "Document"})], "text/markdown" if suffix == ".md" else "text/plain"
        )
    if suffix == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ExtractionError("The uploaded file is not a PDF")
        from pypdf import PdfReader, overwrite_configuration

        overwrite_configuration(
            maximum_declared_stream_length=MAX_UPLOAD_BYTES,
            array_based_stream_maximum_output_length=8 * 1024 * 1024,
            zlib_maximum_output_length=8 * 1024 * 1024,
            lzw_maximum_output_length=8 * 1024 * 1024,
            run_length_maximum_output_length=8 * 1024 * 1024,
            image_maximum_buffer_size=8 * 1024 * 1024,
            page_tree_maximum_entries=1000,
        )

        try:
            reader = PdfReader(io.BytesIO(data), strict=True)
            if reader.is_encrypted:
                raise ExtractionError("Password-protected PDFs are not supported")
            if len(reader.pages) > MAX_PAGES:
                raise ExtractionError("PDF exceeds the 200-page limit")
            parts: list[tuple[str, dict]] = []
            characters = 0
            for i, page in enumerate(reader.pages):
                value = page.extract_text() or ""
                characters += len(value)
                if characters > MAX_CHARACTERS:
                    raise ExtractionError(
                        "Extracted document exceeds the one-million-character limit"
                    )
                parts.append((value, {"page": i + 1}))
            if not any(text.strip() for text, _ in parts):
                raise ExtractionError(
                    "This PDF has no extractable text. OCR is needed for scanned documents"
                )
            return _join(parts, "application/pdf")
        except ExtractionError:
            raise
        except Exception as exc:
            raise ExtractionError(
                "The PDF could not be read; check that it is a valid, unencrypted PDF"
            ) from exc
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ExtractionError("The uploaded file is not a DOCX document")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES:
                raise ExtractionError("DOCX archive expands beyond the supported size")
            if any(item.flag_bits & 1 for item in entries):
                raise ExtractionError("Encrypted DOCX archives are not supported")
            if "word/document.xml" not in archive.namelist():
                raise ExtractionError("The archive does not contain a DOCX document")
            # Never follow external relationships (including attached templates and images).
            for name in archive.namelist():
                if name.endswith(".rels"):
                    relations = ElementTree.fromstring(archive.read(name))
                    for relation in relations:
                        if relation.attrib.get(
                            "TargetMode"
                        ) == "External" and "hyperlink" not in relation.attrib.get("Type", ""):
                            raise ExtractionError(
                                "DOCX external resource references are not supported"
                            )
        from docx import Document

        document = Document(io.BytesIO(data))
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        parts = []
        paragraph_number = table_number = 0
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                paragraph_number += 1
                parts.append((block.text, {"section": f"Paragraph {paragraph_number}"}))
            elif isinstance(block, Table):
                table_number += 1
                parts.append(
                    (
                        "\n".join(
                            " | ".join(cell.text for cell in row.cells) for row in block.rows
                        ),
                        {"section": f"Table {table_number}"},
                    )
                )
        return _join(
            parts, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("The DOCX file could not be read") from exc


async def extract_upload(filename: str, data: bytes) -> ExtractedDocument:
    """Run parsers in a disposable process; terminate malformed/slow documents."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise ExtractionError("Upload a file no larger than 20 MB")
    from atlas.upload_safety import scan_upload

    await scan_upload(filename, data)
    sandbox = tempfile.TemporaryDirectory(prefix="atlas-parser-")
    # No inherited credentials, proxy settings, home directory, or project .env.
    # An absolute module path works for editable installs without PYTHONPATH.
    process = await asyncio.create_subprocess_exec(
        sys.executable,
        "-I",
        "-B",
        WORKER_SCRIPT,
        "--worker",
        cwd=sandbox.name,
        env={
            "PATH": os.defpath,
            "LANG": "C.UTF-8",
            "HOME": sandbox.name,
            "MAX_FILE_BYTES": str(MAX_UPLOAD_BYTES),
            "MAX_DOCUMENT_PAGES": str(MAX_PAGES),
        },
        stdin=asyncio.subprocess.PIPE,
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.DEVNULL,
    )
    try:
        payload = json.dumps(
            {"filename": Path(filename).name, "data": base64.b64encode(data).decode()}
        ).encode()
        stdout, _ = await asyncio.wait_for(process.communicate(payload), timeout=20)
    except BaseException:
        if process.returncode is None:
            process.kill()
            await process.wait()
        raise
    finally:
        sandbox.cleanup()
    if process.returncode:
        raise ExtractionError("Document extraction failed or exceeded resource limits")
    result = json.loads(stdout)
    if "error" in result:
        raise ExtractionError(result["error"])
    return ExtractedDocument(**result)


if __name__ == "__main__":
    # CPU time bounds protect the host independently of the parent's wall-clock timeout.
    import resource

    resource.setrlimit(resource.RLIMIT_CPU, (15, 16))
    resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
    resource.setrlimit(resource.RLIMIT_FSIZE, (0, 0))
    resource.setrlimit(resource.RLIMIT_NOFILE, (64, 64))
    if sys.platform != "darwin":
        resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024, 768 * 1024 * 1024))

    # CPython audit hooks prevent Python-based parsers from opening network sockets
    # or spawning children. Production also runs this process in the container's
    # non-root, read-only filesystem; this hook is defense in depth, not an OS sandbox.
    parser_read_roots = tuple(
        str(item.resolve()) + os.sep
        for item in (Path(sys.prefix), Path(sys.base_prefix), Path(__file__).parent)
    )

    def parser_audit(event, args):
        if event in {
            "socket.connect",
            "socket.bind",
            "socket.getaddrinfo",
            "subprocess.Popen",
            "os.system",
            "os.posix_spawn",
        }:
            raise PermissionError("Parser network and child processes are disabled")
        if event == "open" and isinstance(args[0], (str, bytes)):
            target = os.path.realpath(os.fsdecode(args[0]))
            if not target.startswith(parser_read_roots):
                raise PermissionError("Parser access outside runtime libraries is disabled")
            if isinstance(args[2], int) and args[2] & (os.O_WRONLY | os.O_RDWR | os.O_CREAT):
                raise PermissionError("Parser filesystem writes are disabled")

    sys.addaudithook(parser_audit)
    try:
        body = json.loads(sys.stdin.buffer.read(MAX_UPLOAD_BYTES * 2))
        result = asdict(
            extract_bytes(body["filename"], base64.b64decode(body["data"], validate=True))
        )
    except Exception as exc:
        result = {
            "error": str(exc) if isinstance(exc, ExtractionError) else "Document extraction failed"
        }
    print(json.dumps(result))
