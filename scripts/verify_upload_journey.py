"""Opt-in live HTTP upload journey; real mail, parser, worker and embedding model.

Run after the local API, worker, Mailpit and models are ready. Creates one clearly
named disposable user/company and preserves the evidence for review. No SQL, mocks,
API keys or paid model calls are used. Authentication secrets are never persisted.
"""

import hashlib
import io
import json
import re
import secrets
import time
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlsplit
from uuid import uuid4

import httpx
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

BASE = "http://127.0.0.1:8100"
MAIL = "http://127.0.0.1:58025"
OUTPUT = Path("artifacts/upgrade/upload-journey.json")


def pdf_bytes(lines):
    writer = PdfWriter()
    for text in lines:
        page = writer.add_blank_page(width=600, height=800)
        if not text:
            continue
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 40 700 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def docx_bytes():
    document = Document()
    document.add_heading("Juniper support handover", 0)
    document.add_paragraph("The handover owner is the Platform team. Escalation code: MAPLE-392.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Deployment"
    table.cell(0, 1).text = "Staging has two replicas"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def checked(response, status=200):
    assert response.status_code == status, (
        f"{response.request.method} {response.request.url.path}: expected {status}, got {response.status_code}"
    )
    return response.json()


def verification_token(client, email):
    deadline = time.monotonic() + 25
    while time.monotonic() < deadline:
        messages = checked(client.get(f"{MAIL}/api/v1/messages"))["messages"]
        for message in messages:
            if not any(recipient["Address"] == email for recipient in message.get("To", [])):
                continue
            content = checked(client.get(f"{MAIL}/api/v1/message/{message['ID']}"))
            match = re.search(r"http://[^\s]+/verify-email[^\s]*", content.get("Text", ""))
            if match:
                return parse_qs(urlsplit(match[0]).query)["token"][0]
        time.sleep(1)
    raise AssertionError("Verification email was not delivered to local Mailpit")


def main():
    started = time.monotonic()
    suffix = uuid4().hex[:12]
    email = f"upload-journey-{suffix}@example.test"
    password = "Local upload verification " + secrets.token_urlsafe(24)  # pragma: allowlist secret
    evidence: dict[str, Any] = {
        "started_at": datetime.now(UTC).isoformat(),
        "checks": [],
        "generation_calls": 0,
    }
    with httpx.Client(
        base_url=BASE, timeout=40, headers={"X-Atlas-Client": "console", "Origin": BASE}
    ) as client:
        checked(
            client.post(
                "/api/auth/register",
                json={
                    "email": email,
                    "name": "Disposable Upload Verification",
                    "password": password,
                },
            ),
            202,
        )
        checked(client.post("/api/auth/verify", json={"token": verification_token(client, email)}))
        checked(client.post("/api/auth/login", json={"email": email, "password": password}))
        organization = checked(
            client.post(
                "/api/organizations",
                json={
                    "name": "Disposable Upload Verification",
                    "slug": f"upload-verification-{suffix}",
                },
            ),
            201,
        )["organization"]
        client.headers["X-Atlas-Tenant"] = organization["id"]
        space = checked(
            client.post(
                "/api/spaces", json={"name": "Upload test evidence", "visibility": "restricted"}
            ),
            201,
        )
        evidence.update(tenant_id=organization["id"], space_id=space["id"])
        evidence["checks"].append(
            "Registration, Mailpit verification, login, company and restricted space creation"
        )

        def upload(name, content, media, replace=None):
            data = {"space_id": space["id"]}
            if replace:
                data["replace_document_id"] = replace
            return checked(
                client.post(
                    "/api/library/upload", data=data, files={"file": (name, content, media)}
                ),
                202,
            )

        def ready(record):
            deadline = time.monotonic() + 180
            while time.monotonic() < deadline:
                detail = checked(client.get(f"/api/library/{record['id']}"))
                assert detail.get("pending_status") != "failed", "Real worker failed indexing"
                if (
                    detail["status"] == "ready"
                    and detail["current_version_id"] == record["version_id"]
                ):
                    return detail
                time.sleep(2)
            raise AssertionError(
                "Real worker did not publish the expected version within 180 seconds"
            )

        def version(record):
            return checked(
                client.get(f"/api/library/{record['id']}/versions/{record['version_id']}")
            )

        def download(record, expected):
            result = client.get(
                f"/api/library/{record['id']}/versions/{record['version_id']}/download"
            )
            assert result.status_code == 200 and result.content == expected
            return hashlib.sha256(result.content).hexdigest()

        pdf = pdf_bytes(
            [
                "Juniper production releases require Platform approval code CEDAR-621.",
                "Juniper rollback restores the previous healthy deployment after approval.",
            ]
        )
        pdf_record = upload("two-page-policy.pdf", pdf, "application/pdf")
        ready(pdf_record)
        first_version = version(pdf_record)
        assert [segment["page"] for segment in first_version["source_segments"]] == [1, 2]
        for segment in first_version["source_segments"]:
            assert first_version["content"][segment["start"] : segment["end"]].startswith("Juniper")
        pdf_hash = download(pdf_record, pdf)
        evidence["checks"].append(
            "Real two-page PDF multipart upload, worker indexing, exact page offsets and original byte-for-byte download"
        )

        word = docx_bytes()
        word_record = upload(
            "support-handover.docx",
            word,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        ready(word_record)
        word_version = version(word_record)
        assert "MAPLE-392" in word_version["content"]
        assert "Deployment | Staging has two replicas" in word_version["content"]
        assert any(
            segment.get("section", "").startswith("Paragraph")
            for segment in word_version["source_segments"]
        )
        assert any(
            segment.get("section") == "Table 1" for segment in word_version["source_segments"]
        )
        word_hash = download(word_record, word)
        evidence["checks"].append(
            "Real DOCX multipart upload, worker indexing, paragraph/table references and original byte-for-byte download"
        )

        replacement = pdf_bytes(
            ["Juniper release approval now uses code CEDAR-922 and requires two reviewers."]
        )
        replaced = upload(
            "policy-revision-two.pdf", replacement, "application/pdf", pdf_record["id"]
        )
        assert (
            replaced["id"] == pdf_record["id"]
            and replaced["version_id"] != pdf_record["version_id"]
        )
        download(pdf_record, pdf)
        ready(replaced)
        assert "CEDAR-922" in version(replaced)["content"]
        assert version(pdf_record)["content"] == first_version["content"]
        download(pdf_record, pdf)
        replacement_hash = download(replaced, replacement)
        evidence["checks"].append(
            "Replacement publishes a distinct revision while historical preview/download remains immutable"
        )

        for filename, content, expected in [
            ("broken.pdf", b"this is not a PDF", "PDF"),
            ("scanned.pdf", pdf_bytes([""]), "OCR"),
        ]:
            failure = client.post(
                "/api/library/upload",
                data={"space_id": space["id"]},
                files={"file": (filename, content, "application/pdf")},
            )
            body = checked(failure, 422)
            assert expected.lower() in str(body["detail"]).lower()
        evidence["checks"].append("Malformed and image-only PDFs return actionable 422 errors")
        with httpx.Client(base_url=BASE, timeout=20) as anonymous:
            denied = anonymous.get(
                f"/api/library/{pdf_record['id']}/versions/{pdf_record['version_id']}/download"
            )
            assert denied.status_code == 401
        evidence["checks"].append("Original download rejects unauthenticated requests")
        evidence.update(
            documents={"pdf": pdf_record["id"], "docx": word_record["id"]},
            versions={
                "pdf_original": pdf_record["version_id"],
                "pdf_replacement": replaced["version_id"],
                "docx": word_record["version_id"],
            },
            sha256={
                "pdf_original": pdf_hash,
                "pdf_replacement": replacement_hash,
                "docx": word_hash,
            },
            elapsed_seconds=round(time.monotonic() - started, 2),
            result="passed",
        )
        checked(client.post("/api/auth/logout"))
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(evidence, indent=2) + "\n")
    print(
        json.dumps(
            {
                "result": "passed",
                "checks": len(evidence["checks"]),
                "elapsed_seconds": evidence["elapsed_seconds"],
                "artifact": str(OUTPUT),
            }
        )
    )


if __name__ == "__main__":
    main()
