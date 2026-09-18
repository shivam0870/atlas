import io
import zipfile

import pytest

from atlas.extraction import ExtractionError, extract_bytes, extract_upload


def test_text_extraction_preserves_offsets_and_rejects_binary():
    result = extract_bytes("Guide.MD", "\ufeff Hello 🪴\n\nSecond paragraph ".encode())
    assert result.content == "Hello 🪴\n\nSecond paragraph"
    assert result.content[result.segments[0]["start"] : result.segments[0]["end"]] == result.content
    assert result.media_type == "text/markdown"
    for name, value in [
        ("image.png", b"png"),
        ("fake.txt", b"\x00abc"),
        ("bad.txt", b"\xff"),
        ("empty.txt", b""),
        ("fake.pdf", b"text"),
    ]:
        with pytest.raises(ExtractionError):
            extract_bytes(name, value)


def test_docx_paragraph_table_and_archive_limit():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Release procedure")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Platform"
    output = io.BytesIO()
    doc.save(output)
    result = extract_bytes("runbook.docx", output.getvalue())
    assert "Owner | Platform" in result.content
    assert result.segments[-1]["section"] == "Table 1"
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b" " * (21 * 1024 * 1024))
    with pytest.raises(ExtractionError, match="expands"):
        extract_bytes("bomb.docx", bomb.getvalue())


def test_scanned_pdf_has_actionable_error():
    from pypdf import PdfWriter

    pdf = PdfWriter()
    pdf.add_blank_page(width=100, height=100)
    output = io.BytesIO()
    pdf.write(output)
    with pytest.raises(ExtractionError, match="OCR"):
        extract_bytes("scan.pdf", output.getvalue())


async def test_extraction_uses_disposable_worker():
    result = await extract_upload("hello.txt", b"Bounded parser works")
    assert result.content == "Bounded parser works"


def test_pdf_text_has_stable_page_references():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    for text in ["First page release policy", "Second page rollback policy"]:
        page = writer.add_blank_page(width=600, height=800)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 40 700 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    extracted = extract_bytes("policies.pdf", output.getvalue())
    assert [s["page"] for s in extracted.segments] == [1, 2]
    assert (
        extracted.content[extracted.segments[1]["start"] : extracted.segments[1]["end"]]
        == "Second page rollback policy"
    )


def test_external_docx_template_is_rejected_without_network_access():
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w") as archive:
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr(
            "word/_rels/document.xml.rels",
            '<Relationships><Relationship TargetMode="External" Type="attachedTemplate" Target="https://example.test/private"/></Relationships>',
        )
    with pytest.raises(ExtractionError, match="external resource"):
        extract_bytes("external.docx", payload.getvalue())
